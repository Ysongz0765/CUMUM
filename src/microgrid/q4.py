"""Shared Question 4 data, caching, summaries, and official-workbook export."""
from __future__ import annotations

from copy import copy
from hashlib import sha256
import json
from pathlib import Path
import shutil

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from .execution import execute_fixed_grid_plan
from .io import read_attachment4_prices
from .models.deterministic_dispatch import solve_deterministic_dispatch
from .models.stochastic_dispatch import empirical_cvar
from .prediction import ForecastChoice, build_oos_forecasts, forecast_day


FORMAL_DATES = pd.date_range("2025-02-01", "2025-12-31")
KEY_DATES = pd.to_datetime(["2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21"])


def dependency_signature(paths, extra=None) -> str:
    digest = sha256()
    for path in paths:
        p = Path(path)
        digest.update(str(p.name).encode("utf-8"))
        digest.update(p.read_bytes())
    if extra is not None:
        digest.update(json.dumps(extra, sort_keys=True, ensure_ascii=False).encode("utf-8"))
    return digest.hexdigest()


def price_lookup(raw_dir: Path) -> dict[pd.Timestamp, np.ndarray]:
    prices = read_attachment4_prices(raw_dir)
    return {
        pd.Timestamp(date): day.sort_values("slot").price_yuan_per_kwh.to_numpy(float)
        for date, day in prices.groupby("date", sort=True)
    }


def get_or_build_oos(root: Path, actual: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    """Build a price-independent OOS forecast cache with a formal dependency manifest."""
    raw = root / "data/raw"
    output = root / "outputs/q4"
    output.mkdir(parents=True, exist_ok=True)
    dependencies = [
        root / "src/microgrid/prediction.py",
        root / "src/microgrid/features.py",
        root / "src/microgrid/io.py",
        root / "configs/config.yaml",
        raw / "附件2.xlsx",
    ]
    signature = dependency_signature(dependencies, {"choice": ["ridge", "ridge"]})
    cache_dir = output / "cache" / "forecasts" / signature[:16]
    cache_dir.mkdir(parents=True, exist_ok=True)
    forecast_path = cache_dir / "q4_oos_forecasts.csv"
    if forecast_path.exists():
        forecasts = pd.read_csv(forecast_path, parse_dates=["date"])
    else:
        forecasts = build_oos_forecasts(
            actual,
            pd.date_range("2025-01-08", "2025-12-31"),
            ForecastChoice("ridge", "ridge"),
        )
        forecasts.to_csv(forecast_path, index=False)
    expected = len(pd.date_range("2025-01-08", "2025-12-31")) * 144
    if len(forecasts) != expected or forecasts.duplicated(["date", "slot"]).any():
        raise ValueError("Q4 OOS forecast cache is incomplete or has duplicate keys")
    manifest = {
        "prediction_signature": signature,
        "price_independent": True,
        "forecast_choice": {"load_method": "ridge", "pv_method": "ridge"},
        "rows": len(forecasts),
        "dependencies": [str(path.relative_to(root)) for path in dependencies],
    }
    (output / "q4_oos_forecasts.manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    forecasts.to_csv(output / "q4_oos_forecasts.csv", index=False)
    return forecasts, manifest


def get_or_build_warmup(
    root: Path,
    actual: pd.DataFrame,
    prices: dict[pd.Timestamp, np.ndarray],
    battery: dict,
) -> tuple[pd.DataFrame, dict]:
    """Run the separate January initialization with Attachment 4 prices."""
    output = root / "outputs/q4"
    dependencies = [
        root / "src/microgrid/q4.py",
        root / "src/microgrid/execution.py",
        root / "src/microgrid/models/deterministic_dispatch.py",
        root / "src/microgrid/prediction.py",
        root / "configs/config.yaml",
        root / "data/raw/附件2.xlsx",
        root / "data/raw/附件4.xlsx",
    ]
    signature = dependency_signature(dependencies, {"warmup": "d1_d1_after_day1"})
    cache_dir = output / "cache" / "warmup" / signature[:16]
    cache_dir.mkdir(parents=True, exist_ok=True)
    cache_path = cache_dir / "q4_january_initialization.csv"
    if cache_path.exists():
        warmup = pd.read_csv(cache_path, parse_dates=["date"])
    else:
        rows = []
        soc = float(battery["initial_soc_kwh"])
        for date in pd.date_range("2025-01-01", "2025-01-31"):
            day = actual[actual.date == date].sort_values("slot")
            load_actual = day.load_kwh.to_numpy(float)
            pv_actual = day.pv_actual_kwh.to_numpy(float)
            if date == pd.Timestamp("2025-01-01"):
                load_pred, pv_pred = load_actual, pv_actual
            else:
                point = forecast_day(actual, date, "d1", "d1")
                load_pred = point.load_pred_kwh.to_numpy(float)
                pv_pred = point.pv_pred_kwh.to_numpy(float)
            price = prices[date]
            plan = solve_deterministic_dispatch(load_pred, pv_pred, price, soc, None, battery)
            executed = execute_fixed_grid_plan(
                date, plan.grid_purchase_kwh, load_actual, pv_actual,
                load_pred, pv_pred, price, soc, battery, "q4_warmup"
            )
            rows.append(executed)
            soc = float(executed.soc_end_kwh.iloc[-1])
        warmup = pd.concat(rows, ignore_index=True)
        warmup.to_csv(cache_path, index=False)
    if len(warmup) != 31 * 144:
        raise ValueError(f"Q4 warm-up must have 4464 rows, got {len(warmup)}")
    manifest = {
        "warmup_signature": signature,
        "rows": len(warmup),
        "initial_soc_kwh": float(battery["initial_soc_kwh"]),
        "feb1_initial_soc_kwh": float(warmup.soc_end_kwh.iloc[-1]),
        "cost_excluded_from_formal_evaluation": True,
        "assumption": "Jan 1 uses the observed trajectory only for offline initialization; Jan 2-31 use yesterday-same-slot forecasts and causal execution.",
    }
    warmup.to_csv(output / "q4_january_initialization.csv", index=False)
    (output / "q4_january_initialization.manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return warmup, manifest


def summarize_q4_2(execution: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    daily = execution.groupby(["policy", "date"]).agg(
        plan_cost_yuan=("plan_cost_yuan", "sum"),
        emergency_cost_yuan=("emergency_cost_yuan", "sum"),
        total_cost_yuan=("total_cost_yuan", "sum"),
        emergency_purchase_kwh=("emergency_purchase_kwh", "sum"),
        remaining_energy_kwh=("remaining_energy_kwh", "sum"),
        charge_kwh=("charge_kwh", "sum"),
        discharge_kwh=("discharge_kwh", "sum"),
        soc_min_kwh=("soc_end_kwh", "min"),
        soc_max_kwh=("soc_end_kwh", "max"),
        terminal_soc_kwh=("soc_end_kwh", "last"),
    ).reset_index()
    rows = []
    for policy, group in daily.groupby("policy"):
        weights = np.ones(len(group)) / len(group)
        rows.append({
            "branch": "Q4-2",
            "policy": policy,
            "plan_cost_yuan": group.plan_cost_yuan.sum(),
            "non_emergency_cost_yuan": group.plan_cost_yuan.sum(),
            "emergency_cost_yuan": group.emergency_cost_yuan.sum(),
            "total_cost_yuan": group.total_cost_yuan.sum(),
            "emergency_purchase_kwh": group.emergency_purchase_kwh.sum(),
            "emergency_days": int((group.emergency_purchase_kwh > 1e-8).sum()),
            "daily_cost_std_yuan": group.total_cost_yuan.std(),
            "daily_emergency_cvar95_yuan": empirical_cvar(group.emergency_cost_yuan, weights, .95),
            "daily_total_cvar95_yuan": empirical_cvar(group.total_cost_yuan, weights, .95),
            "remaining_energy_kwh": group.remaining_energy_kwh.sum(),
            "charge_kwh": group.charge_kwh.sum(),
            "discharge_kwh": group.discharge_kwh.sum(),
            "revision_count": 0,
            "min_soc_kwh": group.soc_min_kwh.min(),
            "max_soc_kwh": group.soc_max_kwh.max(),
            "year_end_soc_kwh": group.terminal_soc_kwh.iloc[-1],
        })
    return daily, pd.DataFrame(rows)


def summarize_q4_3(execution: pd.DataFrame, revisions: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    daily = execution.groupby(["policy", "date"]).agg(
        original_plan_cost_yuan=("original_plan_cost_yuan", "sum"),
        final_non_emergency_cost_yuan=("final_non_emergency_cost_yuan", "sum"),
        adjustment_net_cost_yuan=("adjustment_net_cost_yuan", "sum"),
        emergency_cost_yuan=("emergency_cost_yuan", "sum"),
        total_cost_yuan=("total_cost_yuan", "sum"),
        emergency_purchase_kwh=("emergency_purchase_kwh", "sum"),
        remaining_energy_kwh=("remaining_energy_kwh", "sum"),
        charge_kwh=("charge_kwh", "sum"),
        discharge_kwh=("discharge_kwh", "sum"),
        soc_min_kwh=("soc_end_kwh", "min"),
        soc_max_kwh=("soc_end_kwh", "max"),
        terminal_soc_kwh=("soc_end_kwh", "last"),
    ).reset_index()
    rows = []
    for policy, group in daily.groupby("policy"):
        policy_revisions = revisions[revisions.policy == policy] if len(revisions) else revisions
        weights = np.ones(len(group)) / len(group)
        rows.append({
            "branch": "Q4-3",
            "policy": policy,
            "plan_cost_yuan": group.original_plan_cost_yuan.sum(),
            "non_emergency_cost_yuan": group.final_non_emergency_cost_yuan.sum(),
            "emergency_cost_yuan": group.emergency_cost_yuan.sum(),
            "total_cost_yuan": group.total_cost_yuan.sum(),
            "emergency_purchase_kwh": group.emergency_purchase_kwh.sum(),
            "emergency_days": int((group.emergency_purchase_kwh > 1e-8).sum()),
            "daily_cost_std_yuan": group.total_cost_yuan.std(),
            "daily_emergency_cvar95_yuan": empirical_cvar(group.emergency_cost_yuan, weights, .95),
            "daily_total_cvar95_yuan": empirical_cvar(group.total_cost_yuan, weights, .95),
            "remaining_energy_kwh": group.remaining_energy_kwh.sum(),
            "charge_kwh": group.charge_kwh.sum(),
            "discharge_kwh": group.discharge_kwh.sum(),
            "revision_count": int(policy_revisions.triggered.sum()) if len(policy_revisions) else 0,
            "min_soc_kwh": group.soc_min_kwh.min(),
            "max_soc_kwh": group.soc_max_kwh.max(),
            "year_end_soc_kwh": group.terminal_soc_kwh.iloc[-1],
        })
    return daily, pd.DataFrame(rows)


def slot_from_template_label(value) -> int:
    left = str(value).split("-")[0].replace("+1", "")
    hour, minute = map(int, left.split(":"))
    return hour * 6 + minute // 10 + 1


def emergency_groups(day: pd.DataFrame):
    indices = np.flatnonzero(day.emergency_purchase_kwh.to_numpy(float) > 1e-8)
    return np.split(indices, np.where(np.diff(indices) > 1)[0] + 1) if len(indices) else []


def emergency_label(group) -> str:
    start_minutes = int(group[0]) * 10
    end_minutes = (int(group[-1]) + 1) * 10
    label = f"{start_minutes // 60}:{start_minutes % 60:02d}-{(end_minutes // 60) % 24}:{end_minutes % 60:02d}"
    return label + ("+1" if end_minutes >= 1440 else "")


def _write_storage_and_emergency(workbook, execution: pd.DataFrame, storage_index: int, emergency_index: int):
    storage = workbook.worksheets[storage_index]
    styles = [[copy(storage.cell(row, col)._style) for col in range(1, 7)] for row in range(2, 8)]
    storage.delete_rows(2, storage.max_row - 1)
    for day_number, date in enumerate(FORMAL_DATES):
        day = execution[execution.date == date].sort_values("slot")
        start = 2 + day_number * 6
        for block in range(6):
            row = start + block
            for col in range(1, 7):
                storage.cell(row, col)._style = copy(styles[block][col - 1])
            storage.cell(row, 1, date.to_pydatetime() if block == 0 else None)
            storage.cell(row, 2, f"{block * 4}:00-{(block + 1) * 4}:00")
            storage.cell(row, 3, float(day.iloc[block * 24:(block + 1) * 24].charge_kwh.sum()))
            storage.cell(row, 4, float(day.iloc[block * 24:(block + 1) * 24].discharge_kwh.sum()))
        storage.cell(start, 5, "0:00")
        storage.cell(start, 6, float(day.soc_start_kwh.iloc[0]))
        storage.cell(start + 1, 5, "24:00")
        storage.cell(start + 1, 6, float(day.soc_end_kwh.iloc[-1]))
    emergency = workbook.worksheets[emergency_index]
    emergency_styles = [copy(emergency.cell(2, col)._style) for col in range(1, 4)]
    emergency.delete_rows(2, emergency.max_row - 1)
    row = 2
    for date, day in execution.groupby("date", sort=True):
        day = day.sort_values("slot")
        for group_number, group in enumerate(emergency_groups(day)):
            for col in range(1, 4):
                emergency.cell(row, col)._style = copy(emergency_styles[col - 1])
            emergency.cell(row, 1, pd.Timestamp(date).to_pydatetime() if group_number == 0 else None)
            emergency.cell(row, 2, emergency_label(group))
            emergency.cell(row, 3, float(day.iloc[group].emergency_purchase_kwh.sum()))
            row += 1


def _write_plan_sheet(sheet, plans: pd.DataFrame, value_field: str, cost_field: str, execution: pd.DataFrame):
    slot_columns = {slot_from_template_label(sheet.cell(1, col).value): col for col in range(2, 146)}
    for row, date in enumerate(FORMAL_DATES, start=2):
        day_plan = plans[plans.date == date].set_index("slot")[value_field]
        day_execution = execution[execution.date == date]
        for slot, col in slot_columns.items():
            sheet.cell(row, col, float(day_plan.loc[slot]))
        sheet.cell(row, 146, float(day_plan.sum()))
        sheet.cell(row, 147, float(day_execution[cost_field].sum()))


def export_q4_2(root: Path, execution: pd.DataFrame, plans: pd.DataFrame, main_policy: str) -> Path:
    output = root / "outputs/q4"
    target = output / "result4-2.xlsx"
    shutil.copy2(root / "data/raw/templates/result4-2.xlsx", target)
    workbook = load_workbook(target)
    selected_execution = execution[execution.policy == main_policy].copy()
    selected_plans = plans[plans.policy == main_policy].copy()
    _write_plan_sheet(workbook.worksheets[0], selected_plans, "grid_plan_kwh", "plan_cost_yuan", selected_execution)
    _write_storage_and_emergency(workbook, selected_execution, 1, 2)
    workbook.save(target)
    return target


def export_q4_3(root: Path, execution: pd.DataFrame, main_policy: str) -> Path:
    output = root / "outputs/q4"
    target = output / "result4-3.xlsx"
    shutil.copy2(root / "data/raw/templates/result4-3.xlsx", target)
    workbook = load_workbook(target)
    selected = execution[execution.policy == main_policy].copy()
    original = selected[["date", "slot", "original_grid_plan_kwh"]]
    final = selected[["date", "slot", "final_grid_plan_kwh"]]
    _write_plan_sheet(workbook.worksheets[0], original, "original_grid_plan_kwh", "original_plan_cost_yuan", selected)
    _write_plan_sheet(workbook.worksheets[1], final, "final_grid_plan_kwh", "final_non_emergency_cost_yuan", selected)
    _write_storage_and_emergency(workbook, selected, 2, 3)
    workbook.save(target)
    return target
