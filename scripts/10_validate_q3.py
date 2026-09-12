from pathlib import Path
import json, os, sys
import numpy as np
import pandas as pd
import yaml
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parents[1] / 'src'))
from microgrid.io import read_actual, get_q2_daily_price

ROOT = Path(__file__).parents[1]
OUT = ROOT / 'outputs/q3'
cfg = yaml.safe_load((ROOT / 'configs/config.yaml').read_text(encoding='utf-8'))
q3, battery = cfg['q3'], cfg['battery']
tol, cost_tol = 1e-6, 0.01
x = pd.read_csv(Path(os.environ.get('Q3_DETAIL_PATH', OUT / 'q3_execution_detail.csv')), parse_dates=['date'])
rv = pd.read_csv(OUT / 'q3_revision_log.csv', parse_dates=['date', 'decision_time'])
sn = pd.read_csv(OUT / 'q3_plan_snapshots.csv', parse_dates=['date', 'decision_time'])
src = pd.read_csv(OUT / 'q3_scenario_sources.csv', parse_dates=['source_date', 'target_date', 'training_cutoff', 'decision_time'])
actual = read_actual(ROOT / 'data/raw')
actual['load_kwh'] = actual.load_kw / 6
actual['pv_actual_kwh'] = actual.pv_actual_kw / 6
actual = actual[['date', 'slot', 'load_kwh', 'pv_actual_kwh']]
price = get_q2_daily_price(ROOT / 'data/raw').set_index('slot').price_yuan_per_kwh
dates = pd.date_range('2025-02-01', '2025-12-31')
e = {'rows': len(x), 'expected_rows': 4 * 334 * 144, 'policies': sorted(x.policy.unique()), 'main_policy': q3['main_policy']}
e['row_count_ok'] = len(x) == e['expected_rows']
e['keys_unique'] = not x.duplicated(['policy', 'date', 'slot']).any()
e['finite'] = bool(np.isfinite(x.select_dtypes(include=[np.number])).all().all())
e['complete_slots_dates'] = bool(x.groupby(['policy', 'date']).slot.apply(lambda s: np.array_equal(s.to_numpy(), np.arange(1, 145))).all() and all(pd.DatetimeIndex(g.date.unique()).equals(dates) for _, g in x.groupby('policy')))
j = x.merge(actual, on=['date', 'slot'], suffixes=('', '_raw'), validate='many_to_one')
e['actual_matches_attachment2'] = bool(np.allclose(j.load_actual_kwh, j.load_kwh, atol=tol, rtol=0) and np.allclose(j.pv_actual_kwh, j.pv_actual_kwh_raw, atol=tol, rtol=0))
e['attachment1_price_used'] = bool(np.allclose(x.price_yuan_per_kwh, x.slot.map(price), atol=tol, rtol=0))
balance = j.final_grid_plan_kwh + j.emergency_purchase_kwh + j.pv_actual_kwh_raw + j.discharge_kwh - j.load_kwh - j.charge_kwh - j.remaining_energy_kwh
e['max_balance_error_kwh'] = float(balance.abs().max())
e['balance_ok'] = e['max_balance_error_kwh'] <= tol
soc_error = x.soc_end_kwh - (x.soc_start_kwh + battery['charge_efficiency'] * x.charge_kwh - x.discharge_kwh / battery['discharge_efficiency'])
e['max_soc_error_kwh'] = float(soc_error.abs().max())
e['soc_recursion_ok'] = e['max_soc_error_kwh'] <= tol
within = x.groupby(['policy', 'date']).apply(lambda g: np.max(np.abs(g.soc_start_kwh.iloc[1:].to_numpy() - g.soc_end_kwh.iloc[:-1].to_numpy())))
starts = x.groupby(['policy', 'date']).soc_start_kwh.first()
ends = x.groupby(['policy', 'date']).soc_end_kwh.last()
e['max_within_day_soc_gap_kwh'] = float(within.max())
e['max_cross_day_soc_gap_kwh'] = float(max(np.max(np.abs(starts.loc[p].iloc[1:].to_numpy() - ends.loc[p].iloc[:-1].to_numpy())) for p in x.policy.unique()))
e['soc_continuity_ok'] = e['max_within_day_soc_gap_kwh'] <= tol and e['max_cross_day_soc_gap_kwh'] <= tol
e['soc_bounds_ok'] = bool(x.soc_end_kwh.between(battery['soc_min_kwh'] - tol, battery['soc_max_kwh'] + tol).all())
e['charge_limit_ok'] = bool(x.charge_kwh.max() <= battery['max_charge_power_kw'] * cfg['time']['delta_hours'] + tol)
e['discharge_limit_ok'] = bool(x.discharge_kwh.max() <= battery['max_discharge_power_kw'] * cfg['time']['delta_hours'] + tol)
e['nonnegative'] = bool((x[['original_grid_plan_kwh', 'final_grid_plan_kwh', 'emergency_purchase_kwh', 'charge_kwh', 'discharge_kwh', 'remaining_energy_kwh']] >= -tol).all().all())
e['simultaneous_charge_discharge_count'] = int(((x.charge_kwh > tol) & (x.discharge_kwh > tol)).sum())
p = x.slot.map(price).to_numpy()
original_fee = p * x.original_grid_plan_kwh
final_fee = p * x.final_grid_plan_kwh + .5 * p * np.abs(x.final_grid_plan_kwh - x.original_grid_plan_kwh)
emergency_fee = 5 * p * x.emergency_purchase_kwh
e['max_original_fee_error_yuan'] = float(np.max(np.abs(original_fee - x.original_plan_cost_yuan)))
e['max_final_fee_error_yuan'] = float(np.max(np.abs(final_fee - x.final_non_emergency_cost_yuan)))
e['max_emergency_fee_error_yuan'] = float(np.max(np.abs(emergency_fee - x.emergency_cost_yuan)))
e['max_total_fee_error_yuan'] = float(np.max(np.abs(final_fee + emergency_fee - x.total_cost_yuan)))
e['fees_ok'] = max(e['max_original_fee_error_yuan'], e['max_final_fee_error_yuan'], e['max_emergency_fee_error_yuan'], e['max_total_fee_error_yuan']) <= cost_tol
e['scenario_sources_historical'] = bool((src.source_date < src.target_date).all() and (src.training_cutoff < src.target_date).all())
e['scenario_weights_normalized'] = bool(np.allclose(src.groupby(['policy', 'target_date', 'issue_hour', 'decision_time']).weight.sum(), 1, atol=tol, rtol=0))
e['revision_boundaries_ok'] = bool((sn.slot > sn.issue_hour * 6).all())
e['fixed_policies_unchanged'] = bool(np.allclose(x[x.policy.isin(['C0', 'C1'])].original_grid_plan_kwh, x[x.policy.isin(['C0', 'C1'])].final_grid_plan_kwh, atol=tol, rtol=0))

main = x[x.policy == q3['main_policy']]
book = load_workbook(OUT / 'result3.xlsx', data_only=True)
raw = load_workbook(ROOT / 'data/raw/templates/result3.xlsx', data_only=True)
e['template_sheets_preserved'] = book.sheetnames == raw.sheetnames
e['plan_headers_preserved'] = all([book.worksheets[k].cell(1, c).value for c in range(1, 148)] == [raw.worksheets[k].cell(1, c).value for c in range(1, 148)] for k in (0, 1))
plan_ok, block_ok, soc_ok = [], [], []
for rr, date in enumerate(dates, start=2):
    z = main[main.date == date].sort_values('slot')
    for col in range(2, 146):
        left=str(book.worksheets[0].cell(1,col).value).split('-')[0].replace('+1','');hour,minute=map(int,left.split(':'));slot=hour*6+minute//10+1
        row=z[z.slot==slot].iloc[0]
        plan_ok.extend([np.isclose(book.worksheets[0].cell(rr, col).value, row.original_grid_plan_kwh, atol=tol, rtol=0), np.isclose(book.worksheets[1].cell(rr, col).value, row.final_grid_plan_kwh, atol=tol, rtol=0)])
    start = 2 + (rr - 2) * 6
    for k in range(6):
        block_ok.extend([np.isclose(book.worksheets[2].cell(start + k, 3).value, z.charge_kwh.iloc[k * 24:(k + 1) * 24].sum(), atol=tol, rtol=0), np.isclose(book.worksheets[2].cell(start + k, 4).value, z.discharge_kwh.iloc[k * 24:(k + 1) * 24].sum(), atol=tol, rtol=0)])
    soc_ok.extend([np.isclose(book.worksheets[2].cell(start, 6).value, z.soc_start_kwh.iloc[0], atol=tol, rtol=0), np.isclose(book.worksheets[2].cell(start + 1, 6).value, z.soc_end_kwh.iloc[-1], atol=tol, rtol=0)])
e['every_plan_cell_matches'] = bool(all(plan_ok))
e['every_4hour_block_matches'] = bool(all(block_ok))
e['every_daily_soc_matches'] = bool(all(soc_ok))
expected = []
for date, z in main.groupby('date'):
    z = z.sort_values('slot'); ids = np.flatnonzero(z.emergency_purchase_kwh.to_numpy() > 1e-8)
    groups = np.split(ids, np.where(np.diff(ids) > 1)[0] + 1) if len(ids) else []
    for group in groups: expected.append((pd.Timestamp(date), int(group[0]), int(group[-1]), float(z.iloc[group].emergency_purchase_kwh.sum())))
observed, current_date = [], None
for row in range(2, book.worksheets[3].max_row + 1):
    if book.worksheets[3].cell(row, 1).value is not None: current_date = pd.Timestamp(book.worksheets[3].cell(row, 1).value)
    label = str(book.worksheets[3].cell(row, 2).value); left, right = label.replace('+1', '').split('-'); h, m = map(int, left.split(':')); eh, em = map(int, right.split(':'))
    start = h * 6 + m // 10; end = (eh * 60 + em) // 10 - 1
    if '+1' in label or (eh == 0 and em == 0): end = 143
    observed.append((current_date, start, end, float(book.worksheets[3].cell(row, 3).value)))
e['emergency_intervals_exact'] = len(expected) == len(observed) and all(a[:3] == b[:3] and np.isclose(a[3], b[3], atol=tol, rtol=0) for a, b in zip(expected, observed))
required = ['row_count_ok', 'keys_unique', 'finite', 'complete_slots_dates', 'actual_matches_attachment2', 'attachment1_price_used', 'balance_ok', 'soc_recursion_ok', 'soc_continuity_ok', 'soc_bounds_ok', 'charge_limit_ok', 'discharge_limit_ok', 'nonnegative', 'fees_ok', 'scenario_sources_historical', 'scenario_weights_normalized', 'revision_boundaries_ok', 'fixed_policies_unchanged', 'template_sheets_preserved', 'plan_headers_preserved', 'every_plan_cell_matches', 'every_4hour_block_matches', 'every_daily_soc_matches', 'emergency_intervals_exact']
failed = [key for key in required if not e.get(key, False)]
(ROOT / 'reports/q3_acceptance_evidence.json').write_text(json.dumps(e, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(e, ensure_ascii=False, indent=2))
if failed:
    print('FAILED:', failed, file=sys.stderr)
    raise SystemExit(1)
