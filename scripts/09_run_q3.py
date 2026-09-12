from pathlib import Path
from hashlib import sha256
from copy import copy
import json,sys,time,shutil
import numpy as np
import pandas as pd
sys.path.insert(0,str(Path(__file__).parents[1]/'src'))
from microgrid.config import load_config
from microgrid.io import read_processed,read_pv_forecast,get_q2_daily_price
from microgrid.prediction import build_oos_forecasts,ForecastChoice
from microgrid.q3 import run_q3_day
from microgrid.models.stochastic_dispatch import empirical_cvar

ROOT=Path(__file__).parents[1];out=ROOT/'outputs/q3';reports=ROOT/'reports';figs=reports/'figures/q3';out.mkdir(parents=True,exist_ok=True);figs.mkdir(parents=True,exist_ok=True)
cfg=load_config();b=cfg['battery'];q3=cfg['q3'];raw=ROOT/'data/raw';price=get_q2_daily_price(raw).price_yuan_per_kwh.to_numpy();actual=read_processed(ROOT/'data/processed/actual_10min.parquet');forecast_long=read_pv_forecast(raw/'附件3.xlsx')
signature=sha256(b''.join(p.read_bytes() for p in [Path(__file__),ROOT/'src/microgrid/q3.py',ROOT/'src/microgrid/models/adjustment_dispatch.py',ROOT/'src/microgrid/execution.py',ROOT/'src/microgrid/prediction.py',ROOT/'configs/config.yaml',raw/'附件1.xlsx',raw/'附件2.xlsx',raw/'附件3.xlsx'])).hexdigest()
prediction_dependencies=[ROOT/'src/microgrid/prediction.py',ROOT/'src/microgrid/features.py',ROOT/'src/microgrid/io.py',ROOT/'configs/config.yaml',raw/'附件2.xlsx']
prediction_signature=sha256(b''.join(p.read_bytes() for p in prediction_dependencies)).hexdigest()
started=time.time();oos_path=out/'q3_oos_forecasts.csv';oos_manifest=out/'q3_oos_forecasts.manifest.json'
cached_manifest=json.loads(oos_manifest.read_text(encoding='utf-8')) if oos_manifest.exists() else {}
if oos_path.exists() and cached_manifest.get('prediction_signature')==prediction_signature:
    oos=pd.read_csv(oos_path,parse_dates=['date'])
else:
    oos=build_oos_forecasts(actual,pd.date_range('2025-01-08','2025-12-31'),ForecastChoice('ridge','ridge'));oos.to_csv(oos_path,index=False)
    oos_manifest.write_text(json.dumps({'prediction_signature':prediction_signature,'price_independent':True,'dependencies':[str(p.relative_to(ROOT)) for p in prediction_dependencies]},ensure_ascii=False,indent=2),encoding='utf-8')

# January calibration uses only January 22-31 actual execution, starting from the
# already documented Q2 offline warm-up SOC at the end of January 21.
warm=pd.read_csv(ROOT/'outputs/q2/january_initialization.csv',parse_dates=['date']);jan21=float(warm[warm.date==pd.Timestamp('2025-01-21')].soc_end_kwh.iloc[-1])
cal=[]
for delta in q3['delta_candidates_yuan']:
    soc=jan21;costs=[];emergency=[];triggers=0
    for d in pd.date_range(q3['calibration_start'],q3['calibration_end']):
        z=run_q3_day(d,'C3',soc,actual,oos,forecast_long,price,b,q3,float(delta),cfg['random_seed']);soc=float(z.execution.soc_end_kwh.iloc[-1]);costs.append(z.execution.total_cost_yuan.sum());emergency.append(z.execution.emergency_cost_yuan.sum());triggers+=int(z.revisions.triggered.sum())
    cal.append({'delta_yuan':delta,'january_total_cost_yuan':sum(costs),'january_daily_total_cvar95_yuan':empirical_cvar(costs,np.ones(len(costs))/len(costs),.95),'january_daily_emergency_cvar95_yuan':empirical_cvar(emergency,np.ones(len(emergency))/len(emergency),.95),'revision_count':triggers,'terminal_soc_kwh':soc})
cal=pd.DataFrame(cal);best=cal.sort_values(['january_total_cost_yuan','january_daily_emergency_cvar95_yuan','revision_count']).iloc[0];delta=float(best.delta_yuan);cal['selected']=cal.delta_yuan.eq(delta);cal.to_csv(out/'q3_delta_selection.csv',index=False)

common_soc=float(warm.soc_end_kwh.iloc[-1]);socs={p:common_soc for p in ('C0','C1','C2','C3')};execs=[];revisions=[];snapshots=[];sources=[]
dates=pd.date_range('2025-02-01','2025-12-31')
for day_no,d in enumerate(dates,start=1):
    day_results={}
    for policy in ('C0','C1','C2'):
        z=run_q3_day(d,policy,socs[policy],actual,oos,forecast_long,price,b,q3,delta,cfg['random_seed']);day_results[policy]=z;socs[policy]=float(z.execution.soc_end_kwh.iloc[-1])
    if delta==0:
        base=day_results['C2'];ex=base.execution.copy();ex['policy']='C3';rv=base.revisions.copy();rv['policy']='C3';sn=base.snapshots.copy();sn['policy']='C3';ss=base.scenario_sources.copy();ss['policy']='C3';z=type(base)(ex,rv,sn,ss)
    else:z=run_q3_day(d,'C3',socs['C3'],actual,oos,forecast_long,price,b,q3,delta,cfg['random_seed'])
    day_results['C3']=z;socs['C3']=float(z.execution.soc_end_kwh.iloc[-1])
    for z in day_results.values():
        execs.append(z.execution)
        if len(z.revisions):revisions.append(z.revisions)
        if len(z.snapshots):snapshots.append(z.snapshots)
        sources.append(z.scenario_sources)
    if day_no%5==0 or day_no==len(dates):print(f'q3 {day_no}/{len(dates)} elapsed={time.time()-started:.1f}s',flush=True)

execution=pd.concat(execs,ignore_index=True);revision=pd.concat(revisions,ignore_index=True);snapshot=pd.concat(snapshots,ignore_index=True);source=pd.concat(sources,ignore_index=True)
execution.to_csv(out/'q3_execution_detail.csv',index=False);revision.to_csv(out/'q3_revision_log.csv',index=False);snapshot.to_csv(out/'q3_plan_snapshots.csv',index=False);source.to_csv(out/'q3_scenario_sources.csv',index=False)
daily=execution.groupby(['policy','date']).agg(original_plan_cost_yuan=('original_plan_cost_yuan','sum'),final_non_emergency_cost_yuan=('final_non_emergency_cost_yuan','sum'),adjustment_net_cost_yuan=('adjustment_net_cost_yuan','sum'),emergency_cost_yuan=('emergency_cost_yuan','sum'),total_cost_yuan=('total_cost_yuan','sum'),emergency_purchase_kwh=('emergency_purchase_kwh','sum'),remaining_energy_kwh=('remaining_energy_kwh','sum'),charge_kwh=('charge_kwh','sum'),discharge_kwh=('discharge_kwh','sum'),soc_min_kwh=('soc_end_kwh','min'),soc_max_kwh=('soc_end_kwh','max'),terminal_soc_kwh=('soc_end_kwh','last')).reset_index();daily.to_csv(out/'q3_daily_summary.csv',index=False)
comparison=[]
for policy,g in daily.groupby('policy'):
    rv=revision[revision.policy==policy] if len(revision) else pd.DataFrame()
    comparison.append({'policy':policy,'original_plan_cost_yuan':g.original_plan_cost_yuan.sum(),'final_non_emergency_cost_yuan':g.final_non_emergency_cost_yuan.sum(),'adjustment_net_cost_yuan':g.adjustment_net_cost_yuan.sum(),'emergency_cost_yuan':g.emergency_cost_yuan.sum(),'total_cost_yuan':g.total_cost_yuan.sum(),'emergency_purchase_kwh':g.emergency_purchase_kwh.sum(),'daily_emergency_cvar95_yuan':empirical_cvar(g.emergency_cost_yuan,np.ones(len(g))/len(g),.95),'daily_total_cvar95_yuan':empirical_cvar(g.total_cost_yuan,np.ones(len(g))/len(g),.95),'remaining_energy_kwh':g.remaining_energy_kwh.sum(),'revision_count':int(rv.triggered.sum()) if len(rv) else 0,'min_soc_kwh':g.soc_min_kwh.min(),'max_soc_kwh':g.soc_max_kwh.max(),'year_end_soc_kwh':g.terminal_soc_kwh.iloc[-1]})
comparison=pd.DataFrame(comparison);comparison.to_csv(out/'q3_strategy_comparison.csv',index=False)

# Original and latest Attachment 3 forecast accuracy on exactly the same slots.
acc=[]
for d in dates:
    truth=actual[actual.date==d].sort_values('slot').pv_actual_kw.to_numpy()
    from microgrid.q3 import issue_pv_day
    base=issue_pv_day(forecast_long,d,0)*6
    for h in q3['issue_hours']:
        latest=issue_pv_day(forecast_long,d,h)*6;start=h*6;e0=base[start:]-truth[start:];el=latest[start:]-truth[start:]
        acc.append({'date':d,'issue_hour':h,'observations':len(el),'original_mae_kw':np.abs(e0).mean(),'latest_mae_kw':np.abs(el).mean(),'original_rmse_kw':np.sqrt(np.mean(e0**2)),'latest_rmse_kw':np.sqrt(np.mean(el**2))})
pd.DataFrame(acc).to_csv(out/'q3_forecast_update_accuracy.csv',index=False)

# Official workbook: only the preselected main policy.
main=q3['main_policy'];m=execution[execution.policy==main].copy();shutil.copy2(raw/'templates/result3.xlsx',out/'result3.xlsx')
from openpyxl import load_workbook
wb=load_workbook(out/'result3.xlsx')
for ws,field,costfield in [(wb.worksheets[0],'original_grid_plan_kwh','original_plan_cost_yuan'),(wb.worksheets[1],'final_grid_plan_kwh','final_non_emergency_cost_yuan')]:
    def slot_from_label(value):
        left=str(value).split('-')[0].replace('+1','');h,minute=map(int,left.split(':'));return h*6+minute//10+1
    slot_columns={slot_from_label(ws.cell(1,c).value):c for c in range(2,146)}
    for rr,d in enumerate(dates,start=2):
        x=m[m.date==d].sort_values('slot').set_index('slot')
        for slot,cc in slot_columns.items():ws.cell(rr,cc,float(x.loc[slot,field]))
        ws.cell(rr,146,float(x[field].sum()));ws.cell(rr,147,float(x[costfield].sum()))
ws=wb.worksheets[2];styles=[[copy(ws.cell(r,c)._style) for c in range(1,7)] for r in range(2,8)];ws.delete_rows(2,ws.max_row-1)
for di,d in enumerate(dates):
    x=m[m.date==d].sort_values('slot');start=2+di*6
    for j in range(6):
        rr=start+j
        for c in range(1,7):ws.cell(rr,c)._style=copy(styles[j][c-1])
        ws.cell(rr,1,d.to_pydatetime() if j==0 else None);ws.cell(rr,2,f'{j*4}:00-{(j+1)*4}:00');ws.cell(rr,3,float(x.iloc[j*24:(j+1)*24].charge_kwh.sum()));ws.cell(rr,4,float(x.iloc[j*24:(j+1)*24].discharge_kwh.sum()))
    ws.cell(start,5,'0:00');ws.cell(start,6,float(x.soc_start_kwh.iloc[0]));ws.cell(start+1,5,'24:00');ws.cell(start+1,6,float(x.soc_end_kwh.iloc[-1]))
ws=wb.worksheets[3];styles=[copy(ws.cell(2,c)._style) for c in range(1,4)];ws.delete_rows(2,ws.max_row-1);erow=2
for d,x in m.groupby('date'):
    x=x.sort_values('slot');ids=np.flatnonzero(x.emergency_purchase_kwh.to_numpy()>1e-8);groups=np.split(ids,np.where(np.diff(ids)>1)[0]+1) if len(ids) else []
    for gi,g in enumerate(groups):
        for c in range(1,4):ws.cell(erow,c)._style=copy(styles[c-1])
        sm=int(g[0])*10;em=(int(g[-1])+1)*10;label=f'{sm//60}:{sm%60:02d}-{(em//60)%24}:{em%60:02d}'+('+1' if em>=1440 else '')
        ws.cell(erow,1,pd.Timestamp(d).to_pydatetime() if gi==0 else None);ws.cell(erow,2,label);ws.cell(erow,3,float(x.iloc[g].emergency_purchase_kwh.sum()));erow+=1
wb.save(out/'result3.xlsx')

keydates=pd.to_datetime(['2025-03-20','2025-06-21','2025-09-23','2025-12-21']);m[m.date.isin(keydates)].to_csv(out/'q3_key_dates_detail.csv',index=False)
manifest={'signature':signature,'prediction_signature':prediction_signature,'runtime_seconds':time.time()-started,'python':sys.executable,'common_feb1_soc_kwh':common_soc,'selected_delta_yuan':delta,'main_policy':main,'scenario_count':q3['scenario_count'],'rows':len(execution)}
(out/'run_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8');print(comparison.to_string(index=False));print(json.dumps(manifest,ensure_ascii=False,indent=2))
