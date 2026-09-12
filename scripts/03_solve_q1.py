from pathlib import Path
import sys,json,shutil
import numpy as np,pandas as pd
sys.path.insert(0,str(Path(__file__).parents[1]/'src'))
from microgrid.io import read_q1,resolve_raw_file
from microgrid.config import load_config
from microgrid.models.deterministic_dispatch import solve_deterministic_dispatch,validate_dispatch_result
from microgrid.result_template import build_q1_time_mapping
ROOT=Path(__file__).parents[1]; raw=ROOT/'data/raw'; out=ROOT/'outputs/q1'; figs=ROOT/'reports/figures/q1'; out.mkdir(parents=True,exist_ok=True);figs.mkdir(parents=True,exist_ok=True)
conf=load_config(); b=conf['battery']; q=read_q1(resolve_raw_file(raw,'attachment1')); mapping=build_q1_time_mapping(q,raw/'templates/result1.xlsx'); mapping.to_csv(ROOT/'reports/q1_time_mapping.csv',index=False)
r=solve_deterministic_dispatch(q.load_kwh,q.pv_forecast_kwh,q.price_yuan_per_kwh,b['initial_soc_kwh'],b['initial_soc_kwh'],b); chk=validate_dispatch_result(r,q.load_kwh.to_numpy(),q.pv_forecast_kwh.to_numpy(),q.price_yuan_per_kwh.to_numpy(),b['initial_soc_kwh'],b['initial_soc_kwh'],b)
base=np.maximum(q.load_kwh-q.pv_forecast_kwh,0); basecost=float(np.dot(q.price_yuan_per_kwh,base)); df=q.copy();df['grid_purchase_kwh']=r.grid_purchase_kwh;df['charge_kwh']=r.charge_kwh;df['discharge_kwh']=r.discharge_kwh;df['curtailment_kwh']=r.curtailment_kwh;df['soc_start_kwh']=np.r_[b['initial_soc_kwh'],r.soc_kwh[:-1]];df['soc_end_kwh']=r.soc_kwh;df['grid_cost_yuan']=df.price_yuan_per_kwh*df.grid_purchase_kwh;df.to_csv(out/'q1_dispatch.csv',index=False)
summary={'baseline_cost_yuan':basecost,'optimized_cost_yuan':r.total_cost_yuan,'saving_yuan':basecost-r.total_cost_yuan,'saving_rate':(basecost-r.total_cost_yuan)/basecost,'total_load_kwh':float(q.load_kwh.sum()),'total_pv_kwh':float(q.pv_forecast_kwh.sum()),'total_grid_purchase_kwh':float(r.grid_purchase_kwh.sum()),'total_charge_kwh':float(r.charge_kwh.sum()),'total_discharge_kwh':float(r.discharge_kwh.sum()),'total_curtailment_kwh':float(r.curtailment_kwh.sum()),'initial_soc_kwh':b['initial_soc_kwh'],'terminal_soc_kwh':float(r.soc_kwh[-1]),'min_soc_kwh':float(r.soc_kwh.min()),'max_soc_kwh':float(r.soc_kwh.max()),**chk};(out/'q1_summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8');pd.DataFrame([summary]).to_csv(out/'q1_summary.csv',index=False)
selected=df[df.interval_start.dt.hour.isin([10,12,14,16,18,20])&df.interval_start.dt.minute.eq(0)][['slot','interval_start','interval_end','grid_purchase_kwh']];selected.to_csv(out/'q1_table1_selected.csv',index=False)
groups=[]
for i in range(6):
 x=df.iloc[i*24:(i+1)*24];groups.append({'interval':f'{i*4}:00-{(i+1)*4}:00','charge_kwh':x.charge_kwh.sum(),'discharge_kwh':x.discharge_kwh.sum()})
pd.DataFrame(groups).to_csv(out/'q1_table2_4hour.csv',index=False)
from openpyxl import load_workbook
shutil.copy2(raw/'templates/result1.xlsx',out/'result1.xlsx');wb=load_workbook(out/'result1.xlsx');ws=wb['计划购电量']
for m in mapping.itertuples(): ws[m.target_cell]=float(df.loc[df.slot==m.slot,'grid_purchase_kwh'].iloc[0])
ws2=wb['充放电量']
for i,g in enumerate(groups,start=2): ws2.cell(i,2,float(g['charge_kwh']));ws2.cell(i,3,float(g['discharge_kwh']))
ws2['E2']=float(b['initial_soc_kwh']);ws2['E3']=float(r.soc_kwh[-1]);wb.save(out/'result1.xlsx')
import matplotlib.pyplot as plt
x=np.arange(1,145)
def save(name): plt.tight_layout();plt.savefig(figs/name,dpi=300);plt.close()
plt.figure(figsize=(12,5));plt.plot(x,df.load_kwh,label='Load');plt.plot(x,df.pv_forecast_kwh,label='PV');plt.plot(x,df.grid_purchase_kwh,label='Grid');plt.ylabel('kWh / 10 min');plt.xlabel('Slot');plt.legend();save('q1_load_pv_grid.png')
fig,ax=plt.subplots(figsize=(12,5));ax.bar(x,df.charge_kwh,label='Charge');ax.bar(x,-df.discharge_kwh,label='Discharge');ax2=ax.twinx();ax2.plot(x,df.soc_end_kwh,color='black');ax.set_ylabel('kWh / 10 min');ax2.set_ylabel('SOC (kWh)');save('q1_storage_soc.png')
fig,ax=plt.subplots(figsize=(12,5));ax.plot(x,df.price_yuan_per_kwh,color='black');ax2=ax.twinx();ax2.bar(x,df.charge_kwh-df.discharge_kwh,alpha=.5);ax.set_ylabel('yuan / kWh');ax2.set_ylabel('Net charge (kWh)');save('q1_price_dispatch.png')
plt.figure(figsize=(12,5));plt.plot(x,base,label='No battery');plt.plot(x,df.grid_purchase_kwh,label='Optimized');plt.ylabel('kWh / 10 min');plt.xlabel('Slot');plt.legend();save('q1_grid_comparison.png')
(ROOT/'reports/q1_time_alignment_validation.md').write_text('# 问题一时间映射验收\n\n附件1的时间值按区间末端解释：首行0:10对应自然日00:00-00:10，末行0:00+1对应23:50-次日00:00。依据是附件2/4采用相同的10分钟端点列，且题目明确0:00+1为次日0:00。result1模板按区间起点从00:10开始并把00:00-00:10+1放在末行，因此仅输出层循环排列为slot 2..144,1；模型、指定时段抽取和4小时汇总始终采用slot 1..144自然日顺序。另一种“标签为区间起点”的解释会把序列推迟10分钟并跨日，且无法解释末标签作为完整日终点，故不采用。\n',encoding='utf-8')
(ROOT/'reports/q1_analysis.md').write_text(f"# 问题一确定性调度\n\n模型采用G、C、D、W和时段末SOC，单位均为kWh。无储能费用{basecost:.6f}元，优化费用{r.total_cost_yuan:.6f}元，节省{summary['saving_rate']:.4%}。购电{summary['total_grid_purchase_kwh']:.6f}kWh，剩余电量{summary['total_curtailment_kwh']:.6f}kWh，SOC范围{summary['min_soc_kwh']:.2f}-{summary['max_soc_kwh']:.2f}kWh，日末{summary['terminal_soc_kwh']:.2f}kWh。最大平衡误差{chk['max_balance_error']:.3e}，最大SOC递推误差{chk['max_soc_equation_error']:.3e}，同时充放{chk['simultaneous_charge_discharge_count']}个时段。\n",encoding='utf-8')
print(json.dumps(summary,ensure_ascii=False,indent=2))
