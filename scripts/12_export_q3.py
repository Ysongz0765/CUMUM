from pathlib import Path
from copy import copy
import shutil,sys
import numpy as np
import pandas as pd
from openpyxl import load_workbook
sys.path.insert(0,str(Path(__file__).parents[1]/'src'))
from microgrid.config import load_config

ROOT=Path(__file__).parents[1];out=ROOT/'outputs/q3';raw=ROOT/'data/raw';cfg=load_config();main=cfg['q3']['main_policy'];x=pd.read_csv(out/'q3_execution_detail.csv',parse_dates=['date']);m=x[x.policy==main];dates=pd.date_range('2025-02-01','2025-12-31')
shutil.copy2(raw/'templates/result3.xlsx',out/'result3.xlsx');wb=load_workbook(out/'result3.xlsx')
def slot_from_label(value):
    left=str(value).split('-')[0].replace('+1','');hour,minute=map(int,left.split(':'));return hour*6+minute//10+1
for ws,field,costfield in [(wb.worksheets[0],'original_grid_plan_kwh','original_plan_cost_yuan'),(wb.worksheets[1],'final_grid_plan_kwh','final_non_emergency_cost_yuan')]:
    slot_columns={slot_from_label(ws.cell(1,c).value):c for c in range(2,146)}
    for rr,date in enumerate(dates,start=2):
        day=m[m.date==date].set_index('slot')
        for slot,col in slot_columns.items():ws.cell(rr,col,float(day.loc[slot,field]))
        ws.cell(rr,146,float(day[field].sum()));ws.cell(rr,147,float(day[costfield].sum()))
ws=wb.worksheets[2];styles=[[copy(ws.cell(r,c)._style) for c in range(1,7)] for r in range(2,8)];ws.delete_rows(2,ws.max_row-1)
for di,date in enumerate(dates):
    day=m[m.date==date].sort_values('slot');start=2+di*6
    for k in range(6):
        row=start+k
        for c in range(1,7):ws.cell(row,c)._style=copy(styles[k][c-1])
        ws.cell(row,1,date.to_pydatetime() if k==0 else None);ws.cell(row,2,f'{k*4}:00-{(k+1)*4}:00');ws.cell(row,3,float(day.iloc[k*24:(k+1)*24].charge_kwh.sum()));ws.cell(row,4,float(day.iloc[k*24:(k+1)*24].discharge_kwh.sum()))
    ws.cell(start,5,'0:00');ws.cell(start,6,float(day.soc_start_kwh.iloc[0]));ws.cell(start+1,5,'24:00');ws.cell(start+1,6,float(day.soc_end_kwh.iloc[-1]))
ws=wb.worksheets[3];styles=[copy(ws.cell(2,c)._style) for c in range(1,4)];ws.delete_rows(2,ws.max_row-1);row=2
for date,day in m.groupby('date'):
    day=day.sort_values('slot');ids=np.flatnonzero(day.emergency_purchase_kwh.to_numpy()>1e-8);groups=np.split(ids,np.where(np.diff(ids)>1)[0]+1) if len(ids) else []
    for group_no,group in enumerate(groups):
        for c in range(1,4):ws.cell(row,c)._style=copy(styles[c-1])
        sm=int(group[0])*10;em=(int(group[-1])+1)*10;label=f'{sm//60}:{sm%60:02d}-{(em//60)%24}:{em%60:02d}'+('+1' if em>=1440 else '')
        ws.cell(row,1,pd.Timestamp(date).to_pydatetime() if group_no==0 else None);ws.cell(row,2,label);ws.cell(row,3,float(day.iloc[group].emergency_purchase_kwh.sum()));row+=1
wb.save(out/'result3.xlsx');print(out/'result3.xlsx')
