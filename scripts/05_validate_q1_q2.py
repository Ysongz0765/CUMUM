from pathlib import Path
import os
import sys
import json,hashlib,sys
import numpy as np,pandas as pd
from openpyxl import load_workbook
import yaml
sys.path.insert(0,str(Path(__file__).parents[1]/'src'))
from microgrid.io import read_processed,get_q2_daily_price
ROOT=Path(__file__).parents[1]; q1=pd.read_csv(ROOT/'outputs/q1/q1_dispatch.csv',parse_dates=['interval_start','interval_end']); q2=pd.read_csv(Path(os.environ.get('Q2_DETAIL_PATH',ROOT/'outputs/q2/q2_execution_detail.csv')),parse_dates=['date']); plans=pd.read_csv(ROOT/'outputs/q2/q2_day_ahead_plans.csv',parse_dates=['date']); scenarios=pd.read_csv(ROOT/'outputs/q2/scenario_sources.csv',parse_dates=['source_date','target_date','training_cutoff']); warm=pd.read_csv(ROOT/'outputs/q2/january_initialization.csv',parse_dates=['date'])
cfg=yaml.safe_load((ROOT/'configs/config.yaml').read_text(encoding='utf-8')); charge_limit=float(cfg['battery']['max_charge_power_kw'])*float(cfg['time']['delta_hours']); discharge_limit=float(cfg['battery']['max_discharge_power_kw'])*float(cfg['time']['delta_hours'])
evidence={'q1_rows':len(q1),'q2_rows':len(q2),'q2_plan_rows':len(plans),'q2_policies':q2.policy.unique().tolist(),'q2_price_source':'attachment1'}
actual=read_processed(ROOT/'data/processed/actual_10min.parquet')[['date','slot','load_kwh','pv_actual_kwh']]
joined=q2.merge(actual,on=['date','slot'],suffixes=('','_raw'),validate='many_to_one')
evidence['keys_unique']=bool(not q2.duplicated(['policy','date','slot']).any() and not plans.duplicated(['policy','date','slot']).any())
evidence['finite_values']=bool(np.isfinite(q2.select_dtypes(include=[np.number])).all().all())
expected_dates=pd.date_range('2025-02-01','2025-12-31')
evidence['complete_policy_dates']=bool(all(pd.DatetimeIndex(g.date.unique()).equals(expected_dates) for _,g in q2.groupby('policy')))
evidence['actual_matches_attachment2']=bool(np.allclose(joined.load_actual_kwh,joined.load_kwh,atol=1e-6,rtol=0) and np.allclose(joined.pv_actual_kwh_raw,joined.pv_actual_kwh,atol=1e-6,rtol=0))
raw_balance=joined.grid_purchase_kwh+joined.emergency_purchase_kwh+joined.pv_actual_kwh_raw+joined.discharge_kwh-joined.load_kwh-joined.charge_kwh-joined.remaining_energy_kwh
evidence['max_independent_balance_error']=float(raw_balance.abs().max());evidence['independent_balance_within_1e-6']=evidence['max_independent_balance_error']<=1e-6
tariff_raw=get_q2_daily_price(ROOT/'data/raw').set_index('slot').price_yuan_per_kwh
expected_price=q2.slot.map(tariff_raw)
evidence['max_plan_fee_error_yuan']=float((q2.plan_cost_yuan-expected_price*q2.grid_purchase_kwh).abs().max())
evidence['max_emergency_fee_error_yuan']=float((q2.emergency_cost_yuan-5*expected_price*q2.emergency_purchase_kwh).abs().max())
evidence['independent_fees_within_0_01']=bool(evidence['max_plan_fee_error_yuan']<=.01 and evidence['max_emergency_fee_error_yuan']<=.01 and np.allclose(q2.total_cost_yuan,q2.plan_cost_yuan+q2.emergency_cost_yuan,atol=.01,rtol=0))
evidence['q1_slot_sequence']=bool(np.array_equal(q1.slot.to_numpy(),np.arange(1,145)));evidence['q2_slot_sequence']=bool(q2.groupby(['policy','date']).slot.apply(lambda x:np.array_equal(x.to_numpy(),np.arange(1,145))).all());merged=q2.merge(plans,on=['date','slot','policy']);evidence['g_fixed_after_0000']=bool(np.allclose(merged.grid_purchase_kwh,merged.grid_plan_kwh))
evidence['max_execution_balance_error']=float(q2.balance_error_kwh.abs().max()); soc_err=q2.soc_end_kwh-(q2.soc_start_kwh+0.9*q2.charge_kwh-q2.discharge_kwh/0.9);evidence['max_execution_soc_error']=float(soc_err.abs().max());evidence['nonnegative']=bool((q2[['grid_purchase_kwh','emergency_purchase_kwh','charge_kwh','discharge_kwh','remaining_energy_kwh']]>=-1e-8).all().all());evidence['soc_bounds']=bool((q2.soc_end_kwh.between(1200-1e-8,10800+1e-8)).all());evidence['annual_cost_reconciliation']=bool(np.isclose(q2.groupby('policy').total_cost_yuan.sum().to_numpy(),q2.groupby('policy').apply(lambda x:(x.plan_cost_yuan+x.emergency_cost_yuan).sum()).to_numpy()).all());
evidence['warmup_rows']=len(warm);evidence['common_feb1_soc_kwh']=float(warm.soc_end_kwh.iloc[-1]);starts=q2.groupby(['policy','date']).soc_start_kwh.first();ends=q2.groupby(['policy','date']).soc_end_kwh.last();evidence['cross_day_soc_continuity']=bool(all(np.allclose(starts.loc[p].iloc[1:],ends.loc[p].iloc[:-1]) for p in q2.policy.unique()));evidence['scenario_sources_strictly_historical']=bool((scenarios.source_date<scenarios.target_date).all() and (scenarios.training_cutoff<scenarios.target_date).all());evidence['scenario_weights_normalized']=bool(np.allclose(scenarios.groupby('target_date').weight.sum(),1)); tariff=pd.read_csv(ROOT/'outputs/q1/q1_dispatch.csv').sort_values('slot').price_yuan_per_kwh.to_numpy();evidence['attachment1_tariff_used']=bool(all(np.allclose(g.sort_values('slot').price_yuan_per_kwh,tariff) for _,g in q2.groupby(['policy','date'])));evidence['b3_differs_from_b1']=bool(not np.allclose(plans[plans.policy=='B1'].grid_plan_kwh,plans[plans.policy=='B3'].grid_plan_kwh));evidence['b4_differs_from_b3']=bool(not np.allclose(plans[plans.policy=='B3'].grid_plan_kwh,plans[plans.policy=='B4'].grid_plan_kwh));evidence['max_charge_kwh']=float(q2.charge_kwh.max());evidence['max_discharge_kwh']=float(q2.discharge_kwh.max());evidence['simultaneous_charge_discharge_count']=int(((q2.charge_kwh>1e-8)&(q2.discharge_kwh>1e-8)).sum())
for fn in ['result1.xlsx','result2.xlsx']:
    p=ROOT/'outputs'/'q1'/fn if fn=='result1.xlsx' else ROOT/'outputs'/'q2'/fn;w=load_workbook(p,data_only=True);evidence[fn]={'sheets':w.sheetnames,'shapes':{s:[w[s].max_row,w[s].max_column] for s in w.sheetnames}}
evidence['result1_no_F_write']=load_workbook(ROOT/'outputs/q1/result1.xlsx',data_only=True).worksheets[1].cell(2,6).value is None
w1raw=load_workbook(ROOT/'data/raw/templates/result1.xlsx',data_only=True);w1out=load_workbook(ROOT/'outputs/q1/result1.xlsx',data_only=True);evidence['result1_headers_labels_preserved']=bool([w1raw.worksheets[0].cell(r,1).value for r in range(1,146)]==[w1out.worksheets[0].cell(r,1).value for r in range(1,146)] and [w1raw.worksheets[1].cell(r,1).value for r in range(1,8)]==[w1out.worksheets[1].cell(r,1).value for r in range(1,8)])
w1=load_workbook(ROOT/'outputs/q1/result1.xlsx',data_only=True);evidence['result1_charge_sum_matches']=bool(np.isclose(sum(w1.worksheets[1].cell(r,2).value for r in range(2,8)),q1.charge_kwh.sum()));evidence['result1_discharge_sum_matches']=bool(np.isclose(sum(w1.worksheets[1].cell(r,3).value for r in range(2,8)),q1.discharge_kwh.sum()))
w2=load_workbook(ROOT/'outputs/q2/result2.xlsx',data_only=True);main=q2[q2.policy=='B4'];evidence['result2_plan_total_matches']=bool(np.isclose(sum(w2.worksheets[0].cell(r,146).value for r in range(2,336)),main.grid_purchase_kwh.sum()));evidence['result2_plan_cost_matches']=bool(np.isclose(sum(w2.worksheets[0].cell(r,147).value for r in range(2,336)),main.plan_cost_yuan.sum()));evidence['result2_charge_sum_matches']=bool(np.isclose(sum(w2.worksheets[1].cell(r,3).value for r in range(2,2006)),main.charge_kwh.sum()));evidence['result2_discharge_sum_matches']=bool(np.isclose(sum(w2.worksheets[1].cell(r,4).value for r in range(2,2006)),main.discharge_kwh.sum()));evidence['result2_emergency_sum_matches']=bool(np.isclose(sum(w2.worksheets[2].cell(r,3).value for r in range(2,w2.worksheets[2].max_row+1)),main.emergency_purchase_kwh.sum()))
raw2=load_workbook(ROOT/'data/raw/templates/result2.xlsx',data_only=True);evidence['result2_plan_headers_preserved']=bool([raw2.worksheets[0].cell(1,c).value for c in range(1,148)]==[w2.worksheets[0].cell(1,c).value for c in range(1,148)])
def _slot_from_label(v):
    s=str(v).split('-')[0].replace('+1','');h,m=map(int,s.split(':'));return h*6+m//10+1
checks=[]
for rr,d in enumerate(pd.date_range('2025-02-01','2025-12-31'),start=2):
    px=plans[(plans.policy=='B4')&(plans.date==d)].set_index('slot').grid_plan_kwh
    for cc in range(2,146): checks.append(np.isclose(w2.worksheets[0].cell(rr,cc).value,px.loc[_slot_from_label(w2.worksheets[0].cell(1,cc).value)]))
evidence['result2_every_plan_cell_matches']=bool(all(checks))
block_ok=[];soc_book_ok=[]
for di,d in enumerate(expected_dates):
    x=main[main.date==d].sort_values('slot');start=2+di*6
    for j in range(6):
        block_ok.extend([np.isclose(w2.worksheets[1].cell(start+j,3).value,x.iloc[j*24:(j+1)*24].charge_kwh.sum(),atol=1e-6,rtol=0),np.isclose(w2.worksheets[1].cell(start+j,4).value,x.iloc[j*24:(j+1)*24].discharge_kwh.sum(),atol=1e-6,rtol=0)])
    soc_book_ok.extend([np.isclose(w2.worksheets[1].cell(start,6).value,x.soc_start_kwh.iloc[0],atol=1e-6,rtol=0),np.isclose(w2.worksheets[1].cell(start+1,6).value,x.soc_end_kwh.iloc[-1],atol=1e-6,rtol=0)])
evidence['result2_every_4hour_block_matches']=bool(all(block_ok));evidence['result2_every_daily_soc_matches']=bool(all(soc_book_ok))
evidence['balance_error_within_1e-6']=evidence['max_execution_balance_error']<=1e-6
evidence['soc_error_within_1e-6']=evidence['max_execution_soc_error']<=1e-6
evidence['charge_limit_kwh']=charge_limit; evidence['discharge_limit_kwh']=discharge_limit
evidence['charge_limit_within_1e-6']=evidence['max_charge_kwh']<=charge_limit+1e-6
evidence['discharge_limit_within_1e-6']=evidence['max_discharge_kwh']<=discharge_limit+1e-6
evidence['annual_cost_reconciliation_strict']=bool(np.allclose(q2.groupby('policy').total_cost_yuan.sum().to_numpy(),q2.groupby('policy').apply(lambda x:(x.plan_cost_yuan+x.emergency_cost_yuan).sum()).to_numpy(),atol=0.01,rtol=0))
evidence['result2_plan_cost_matches_strict']=bool(np.isclose(sum(w2.worksheets[0].cell(r,147).value for r in range(2,336)),main.plan_cost_yuan.sum(),atol=0.01,rtol=0))
required=['keys_unique','finite_values','complete_policy_dates','actual_matches_attachment2','independent_balance_within_1e-6','independent_fees_within_0_01','q1_slot_sequence','q2_slot_sequence','g_fixed_after_0000','nonnegative','soc_bounds','balance_error_within_1e-6','soc_error_within_1e-6','charge_limit_within_1e-6','discharge_limit_within_1e-6','annual_cost_reconciliation_strict','cross_day_soc_continuity','scenario_sources_strictly_historical','scenario_weights_normalized','attachment1_tariff_used','result1_no_F_write','result1_headers_labels_preserved','result1_charge_sum_matches','result1_discharge_sum_matches','result2_plan_total_matches','result2_plan_cost_matches_strict','result2_charge_sum_matches','result2_discharge_sum_matches','result2_emergency_sum_matches','result2_plan_headers_preserved','result2_every_plan_cell_matches','result2_every_4hour_block_matches','result2_every_daily_soc_matches']
failed=[k for k in required if not evidence.get(k,False)]
(ROOT/'reports/acceptance_evidence.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2,default=str),encoding='utf-8');print(json.dumps(evidence,ensure_ascii=False,indent=2))
if failed: print('FAILED CHECKS:',failed,file=sys.stderr); raise SystemExit(1)
