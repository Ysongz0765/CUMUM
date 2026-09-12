from pathlib import Path
import json
import os
import sys

import numpy as np
import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parents[1] / "src"))
from microgrid.config import load_config
from microgrid.io import read_actual, read_attachment4_prices
from microgrid.q4 import FORMAL_DATES, emergency_groups, emergency_label, slot_from_template_label


ROOT = Path(__file__).parents[1]
OUT = ROOT / "outputs/q4"
TOL = 1e-6
COST_TOL = 0.01


def raw_inputs():
    actual = read_actual(ROOT / "data/raw")
    actual["load_kwh_raw"] = actual.load_kw / 6
    actual["pv_kwh_raw"] = actual.pv_actual_kw / 6
    prices = read_attachment4_prices(ROOT / "data/raw").rename(columns={"price_yuan_per_kwh": "price_raw"})
    return actual[["date", "slot", "load_kwh_raw", "pv_kwh_raw"]].merge(
        prices, on=["date", "slot"], validate="one_to_one"
    )


def physical_checks(frame, expected_policies, battery, prefix, evidence):
    joined = frame.merge(raw_inputs(), on=["date", "slot"], validate="many_to_one")
    evidence[f"{prefix}_rows"] = len(frame)
    evidence[f"{prefix}_row_count_ok"] = len(frame) == len(expected_policies) * 334 * 144
    evidence[f"{prefix}_policies_ok"] = set(frame.policy) == set(expected_policies)
    evidence[f"{prefix}_keys_unique"] = not frame.duplicated(["policy", "date", "slot"]).any()
    evidence[f"{prefix}_finite"] = bool(np.isfinite(frame.select_dtypes(include=[np.number])).all().all())
    complete = frame.groupby(["policy", "date"]).slot.apply(lambda values: np.array_equal(values.to_numpy(), np.arange(1, 145)))
    date_complete = all(pd.DatetimeIndex(group.date.unique()).equals(FORMAL_DATES) for _, group in frame.groupby("policy"))
    evidence[f"{prefix}_complete_days_slots"] = bool(complete.all() and date_complete)
    evidence[f"{prefix}_actual_matches_raw_attachment2"] = bool(
        np.allclose(joined.load_actual_kwh, joined.load_kwh_raw, atol=TOL, rtol=0)
        and np.allclose(joined.pv_actual_kwh, joined.pv_kwh_raw, atol=TOL, rtol=0)
    )
    evidence[f"{prefix}_price_matches_raw_attachment4"] = bool(
        np.allclose(joined.price_yuan_per_kwh, joined.price_raw, atol=TOL, rtol=0)
    )
    grid = joined.grid_purchase_kwh if prefix == "q4_2" else joined.final_grid_plan_kwh
    balance = grid + joined.emergency_purchase_kwh + joined.pv_kwh_raw + joined.discharge_kwh - joined.load_kwh_raw - joined.charge_kwh - joined.remaining_energy_kwh
    soc_error = joined.soc_end_kwh - (
        joined.soc_start_kwh + battery["charge_efficiency"] * joined.charge_kwh
        - joined.discharge_kwh / battery["discharge_efficiency"]
    )
    evidence[f"{prefix}_max_balance_error_kwh"] = float(balance.abs().max())
    evidence[f"{prefix}_max_soc_error_kwh"] = float(soc_error.abs().max())
    evidence[f"{prefix}_balance_ok"] = evidence[f"{prefix}_max_balance_error_kwh"] <= TOL
    evidence[f"{prefix}_soc_recursion_ok"] = evidence[f"{prefix}_max_soc_error_kwh"] <= TOL
    within = frame.groupby(["policy", "date"]).apply(
        lambda group: np.max(np.abs(group.soc_start_kwh.iloc[1:].to_numpy() - group.soc_end_kwh.iloc[:-1].to_numpy())),
        include_groups=False,
    )
    starts = frame.groupby(["policy", "date"]).soc_start_kwh.first()
    ends = frame.groupby(["policy", "date"]).soc_end_kwh.last()
    cross = max(np.max(np.abs(starts.loc[p].iloc[1:].to_numpy() - ends.loc[p].iloc[:-1].to_numpy())) for p in expected_policies)
    evidence[f"{prefix}_max_within_day_soc_gap_kwh"] = float(within.max())
    evidence[f"{prefix}_max_cross_day_soc_gap_kwh"] = float(cross)
    evidence[f"{prefix}_soc_continuity_ok"] = within.max() <= TOL and cross <= TOL
    limit_charge = battery["max_charge_power_kw"] / 6
    limit_discharge = battery["max_discharge_power_kw"] / 6
    evidence[f"{prefix}_soc_bounds_ok"] = bool(frame.soc_end_kwh.between(battery["soc_min_kwh"] - TOL, battery["soc_max_kwh"] + TOL).all())
    evidence[f"{prefix}_charge_limit_ok"] = bool(frame.charge_kwh.max() <= limit_charge + TOL)
    evidence[f"{prefix}_discharge_limit_ok"] = bool(frame.discharge_kwh.max() <= limit_discharge + TOL)
    nonnegative_fields = ["emergency_purchase_kwh", "charge_kwh", "discharge_kwh", "remaining_energy_kwh"]
    nonnegative_fields += ["grid_purchase_kwh"] if prefix == "q4_2" else ["original_grid_plan_kwh", "final_grid_plan_kwh"]
    evidence[f"{prefix}_nonnegative"] = bool((frame[nonnegative_fields] >= -TOL).all().all())
    evidence[f"{prefix}_simultaneous_charge_discharge_count"] = int(((frame.charge_kwh > TOL) & (frame.discharge_kwh > TOL)).sum())
    evidence[f"{prefix}_no_simultaneous_charge_discharge"] = evidence[f"{prefix}_simultaneous_charge_discharge_count"] == 0
    return joined


def workbook_checks(frame, workbook_name, template_name, main_policy, branch, evidence):
    selected = frame[frame.policy == main_policy]
    book = load_workbook(OUT / workbook_name, data_only=True)
    template = load_workbook(ROOT / "data/raw/templates" / template_name, data_only=True)
    evidence[f"{branch}_template_sheets_preserved"] = book.sheetnames == template.sheetnames
    plan_sheet_count = 1 if branch == "q4_2" else 2
    evidence[f"{branch}_plan_headers_preserved"] = all(
        [book.worksheets[i].cell(1, col).value for col in range(1, 148)]
        == [template.worksheets[i].cell(1, col).value for col in range(1, 148)]
        for i in range(plan_sheet_count)
    )
    plan_ok, daily_ok, block_ok, soc_ok = [], [], [], []
    for row, date in enumerate(FORMAL_DATES, start=2):
        day = selected[selected.date == date].sort_values("slot").set_index("slot")
        sheet_specs = [(0, "grid_purchase_kwh", "plan_cost_yuan")] if branch == "q4_2" else [
            (0, "original_grid_plan_kwh", "original_plan_cost_yuan"),
            (1, "final_grid_plan_kwh", "final_non_emergency_cost_yuan"),
        ]
        for sheet_index, field, cost_field in sheet_specs:
            sheet = book.worksheets[sheet_index]
            for col in range(2, 146):
                slot = slot_from_template_label(sheet.cell(1, col).value)
                plan_ok.append(np.isclose(sheet.cell(row, col).value, day.loc[slot, field], atol=TOL, rtol=0))
            daily_ok.extend([
                np.isclose(sheet.cell(row, 146).value, day[field].sum(), atol=TOL, rtol=0),
                np.isclose(sheet.cell(row, 147).value, day[cost_field].sum(), atol=COST_TOL, rtol=0),
            ])
        storage_index = 1 if branch == "q4_2" else 2
        start = 2 + (row - 2) * 6
        day_reset = day.reset_index()
        for block in range(6):
            block_ok.extend([
                np.isclose(book.worksheets[storage_index].cell(start + block, 3).value, day_reset.iloc[block*24:(block+1)*24].charge_kwh.sum(), atol=TOL, rtol=0),
                np.isclose(book.worksheets[storage_index].cell(start + block, 4).value, day_reset.iloc[block*24:(block+1)*24].discharge_kwh.sum(), atol=TOL, rtol=0),
            ])
        soc_ok.extend([
            np.isclose(book.worksheets[storage_index].cell(start, 6).value, day_reset.soc_start_kwh.iloc[0], atol=TOL, rtol=0),
            np.isclose(book.worksheets[storage_index].cell(start + 1, 6).value, day_reset.soc_end_kwh.iloc[-1], atol=TOL, rtol=0),
        ])
    evidence[f"{branch}_every_plan_cell_matches"] = bool(all(plan_ok))
    evidence[f"{branch}_daily_totals_costs_match"] = bool(all(daily_ok))
    evidence[f"{branch}_every_4hour_block_matches"] = bool(all(block_ok))
    evidence[f"{branch}_every_daily_soc_matches"] = bool(all(soc_ok))
    emergency_index = 2 if branch == "q4_2" else 3
    expected = []
    for date, day in selected.groupby("date", sort=True):
        day = day.sort_values("slot")
        for group in emergency_groups(day):
            expected.append((pd.Timestamp(date), emergency_label(group), float(day.iloc[group].emergency_purchase_kwh.sum())))
    observed, current_date = [], None
    sheet = book.worksheets[emergency_index]
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, 1).value is not None:
            current_date = pd.Timestamp(sheet.cell(row, 1).value)
        observed.append((current_date, str(sheet.cell(row, 2).value), float(sheet.cell(row, 3).value)))
    evidence[f"{branch}_emergency_intervals_exact"] = len(expected) == len(observed) and all(
        a[:2] == b[:2] and np.isclose(a[2], b[2], atol=TOL, rtol=0) for a, b in zip(expected, observed)
    )


def main():
    config = load_config(); battery = config["battery"]; q4 = config["q4"]
    evidence = {"energy_soc_tolerance_kwh": TOL, "fee_tolerance_yuan": COST_TOL}
    q42 = pd.read_csv(Path(os.environ.get("Q4_2_DETAIL_PATH", OUT / "q4_2_execution_detail.csv")), parse_dates=["date"])
    q43 = pd.read_csv(Path(os.environ.get("Q4_3_DETAIL_PATH", OUT / "q4_3_execution_detail.csv")), parse_dates=["date"])
    plans42 = pd.read_csv(OUT / "q4_2_day_ahead_plans.csv", parse_dates=["date"])
    joined42 = physical_checks(q42, ["B1", "B3", "B4"], battery, "q4_2", evidence)
    joined43 = physical_checks(q43, ["C0", "C1", "C2", "C3"], battery, "q4_3", evidence)
    raw_price42 = joined42.price_raw.to_numpy()
    plan_fee = raw_price42 * joined42.grid_purchase_kwh
    emergency_fee42 = 5 * raw_price42 * joined42.emergency_purchase_kwh
    evidence["q4_2_max_plan_fee_error_yuan"] = float(np.max(np.abs(plan_fee - joined42.plan_cost_yuan)))
    evidence["q4_2_max_emergency_fee_error_yuan"] = float(np.max(np.abs(emergency_fee42 - joined42.emergency_cost_yuan)))
    evidence["q4_2_max_total_fee_error_yuan"] = float(np.max(np.abs(plan_fee + emergency_fee42 - joined42.total_cost_yuan)))
    evidence["q4_2_fee_rows_ok"] = max(evidence["q4_2_max_plan_fee_error_yuan"], evidence["q4_2_max_emergency_fee_error_yuan"], evidence["q4_2_max_total_fee_error_yuan"]) <= COST_TOL
    evidence["q4_2_annual_fee_error_yuan"] = float(abs((plan_fee + emergency_fee42).sum() - joined42.total_cost_yuan.sum()))
    evidence["q4_2_annual_fee_ok"] = evidence["q4_2_annual_fee_error_yuan"] <= COST_TOL
    plan_join = q42.merge(plans42, on=["policy", "date", "slot"], validate="one_to_one")
    evidence["q4_2_day_ahead_g_matches_execution"] = bool(np.allclose(plan_join.grid_purchase_kwh, plan_join.grid_plan_kwh, atol=TOL, rtol=0))
    raw_price43 = joined43.price_raw.to_numpy()
    original_fee = raw_price43 * joined43.original_grid_plan_kwh
    final_fee = raw_price43 * joined43.final_grid_plan_kwh + .5 * raw_price43 * np.abs(joined43.final_grid_plan_kwh - joined43.original_grid_plan_kwh)
    emergency_fee43 = 5 * raw_price43 * joined43.emergency_purchase_kwh
    for name, expected, observed in (
        ("original", original_fee, joined43.original_plan_cost_yuan),
        ("final", final_fee, joined43.final_non_emergency_cost_yuan),
        ("emergency", emergency_fee43, joined43.emergency_cost_yuan),
        ("total", final_fee + emergency_fee43, joined43.total_cost_yuan),
    ):
        evidence[f"q4_3_max_{name}_fee_error_yuan"] = float(np.max(np.abs(expected - observed)))
    evidence["q4_3_fee_rows_ok"] = max(evidence[key] for key in evidence if key.startswith("q4_3_max_") and key.endswith("fee_error_yuan")) <= COST_TOL
    evidence["q4_3_annual_fee_error_yuan"] = float(abs((final_fee + emergency_fee43).sum() - joined43.total_cost_yuan.sum()))
    evidence["q4_3_annual_fee_ok"] = evidence["q4_3_annual_fee_error_yuan"] <= COST_TOL
    warmup = pd.read_csv(OUT / "q4_january_initialization.csv", parse_dates=["date"])
    warm_joined = warmup.merge(raw_inputs(), on=["date", "slot"], validate="one_to_one")
    warm_balance = warm_joined.grid_purchase_kwh + warm_joined.emergency_purchase_kwh + warm_joined.pv_kwh_raw + warm_joined.discharge_kwh - warm_joined.load_kwh_raw - warm_joined.charge_kwh - warm_joined.remaining_energy_kwh
    warm_soc = warm_joined.soc_end_kwh - (warm_joined.soc_start_kwh + battery["charge_efficiency"]*warm_joined.charge_kwh - warm_joined.discharge_kwh/battery["discharge_efficiency"])
    evidence["warmup_rows_ok"] = len(warmup) == 4464
    evidence["warmup_initial_soc_ok"] = abs(warmup.soc_start_kwh.iloc[0] - 6000) <= TOL
    evidence["warmup_price_attachment4_ok"] = bool(np.allclose(warm_joined.price_yuan_per_kwh, warm_joined.price_raw, atol=TOL, rtol=0))
    evidence["warmup_physics_ok"] = warm_balance.abs().max() <= TOL and warm_soc.abs().max() <= TOL
    common_soc = float(warmup.soc_end_kwh.iloc[-1])
    evidence["common_feb1_soc_kwh"] = common_soc
    starts42 = q42[q42.date == FORMAL_DATES[0]].groupby("policy").soc_start_kwh.first()
    starts43 = q43[q43.date == FORMAL_DATES[0]].groupby("policy").soc_start_kwh.first()
    evidence["all_policies_share_q4_feb1_soc"] = bool(np.allclose(pd.concat([starts42, starts43]), common_soc, atol=TOL, rtol=0))
    sources42 = pd.read_csv(OUT / "q4_2_scenario_sources.csv", parse_dates=["source_date", "target_date", "training_cutoff"])
    sources43 = pd.read_csv(OUT / "q4_3_scenario_sources.csv", parse_dates=["source_date", "target_date", "training_cutoff", "decision_time"])
    evidence["q4_2_scenarios_historical"] = bool((sources42.source_date < sources42.target_date).all() and (sources42.training_cutoff < sources42.target_date).all())
    evidence["q4_2_scenario_weights_normalized"] = bool(np.allclose(sources42.groupby("target_date").weight.sum(), 1, atol=TOL, rtol=0))
    evidence["q4_3_scenarios_historical"] = bool((sources43.source_date < sources43.target_date).all() and (sources43.training_cutoff < sources43.target_date).all())
    evidence["q4_3_scenario_weights_normalized"] = bool(np.allclose(sources43.groupby(["policy", "target_date", "issue_hour", "decision_time"]).weight.sum(), 1, atol=TOL, rtol=0))
    snapshots = pd.read_csv(OUT / "q4_3_plan_snapshots.csv", parse_dates=["date", "decision_time"])
    evidence["q4_3_revision_boundaries_ok"] = bool((snapshots.slot > snapshots.issue_hour * 6).all())
    evidence["q4_3_fixed_policies_unchanged"] = bool(np.allclose(q43[q43.policy.isin(["C0", "C1"])].original_grid_plan_kwh, q43[q43.policy.isin(["C0", "C1"])].final_grid_plan_kwh, atol=TOL, rtol=0))
    workbook_checks(q42, "result4-2.xlsx", "result4-2.xlsx", q4["main_policy_q4_2"], "q4_2", evidence)
    workbook_checks(q43, "result4-3.xlsx", "result4-3.xlsx", q4["main_policy_q4_3"], "q4_3", evidence)
    required = [key for key, value in evidence.items() if isinstance(value, (bool, np.bool_))]
    failed = [key for key in required if not bool(evidence[key])]
    evidence["required_check_count"] = len(required)
    evidence["failed_checks"] = failed
    encoded = json.dumps(
        evidence, ensure_ascii=False, indent=2,
        default=lambda value: value.item() if isinstance(value, np.generic) else str(value),
    )
    (ROOT / "reports/q4_acceptance_evidence.json").write_text(encoded, encoding="utf-8")
    print(encoded)
    if failed:
        print("FAILED CHECKS:", failed, file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except SystemExit:
        raise
    except Exception as error:
        failure = {"exception": type(error).__name__, "message": str(error)}
        (ROOT / "reports/q4_acceptance_evidence.json").write_text(json.dumps(failure, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps(failure, ensure_ascii=False, indent=2), file=sys.stderr)
        raise SystemExit(1)
