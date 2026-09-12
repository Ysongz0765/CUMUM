from pathlib import Path
import numpy as np,pandas as pd
from openpyxl import load_workbook
from microgrid.io import read_processed,get_q2_daily_price,read_attachment4_prices,get_daily_price
from microgrid.prediction import forecast_day
from microgrid.models.stochastic_dispatch import solve_two_stage_dispatch,empirical_cvar
from microgrid.execution import execute_fixed_grid_plan
from microgrid.config import load_config

ROOT=Path(__file__).parents[1]
def test_actual_data_chronological_order():
    d=read_processed(ROOT/'data/processed/actual_10min.csv'); assert d.groupby('date').slot.apply(lambda x:np.array_equal(x.to_numpy(),np.arange(1,145))).all()
def test_load_d1_last_is_yesterday_last_slot():
    d=pd.DataFrame({'date':pd.to_datetime(['2025-01-01']*2+['2025-01-02']*2),'slot':[1,2,1,2],'load_kw':[1,9,3,4]}); from microgrid.features import add_lag_features; x=add_lag_features(d); assert x.iloc[-1].load_d1_last==9
def test_q2_uses_attachment1_price():
    p=get_q2_daily_price(ROOT/'data/raw'); assert len(p)==144 and p.iloc[0].price_yuan_per_kwh==0.4248
def test_attachment4_price_interface_is_complete_and_date_specific():
    p=read_attachment4_prices(ROOT/'data/raw')
    assert len(p)==52560 and p.date.nunique()==365
    assert p.groupby('date').slot.nunique().eq(144).all()
    assert np.isfinite(p.price_yuan_per_kwh).all() and (p.price_yuan_per_kwh>0).all()
    jan1=get_daily_price(ROOT/'data/raw','2025-01-01','attachment4')
    jan2=get_daily_price(ROOT/'data/raw','2025-01-02','attachment4')
    assert len(jan1)==144 and not np.allclose(jan1.price_yuan_per_kwh,jan2.price_yuan_per_kwh)
    assert get_daily_price(ROOT/'data/raw','2025-01-01','attachment1').iloc[0].price_yuan_per_kwh==0.4248
def test_q2_plan_has_three_policies_and_144288_rows():
    x=pd.read_csv(ROOT/'outputs/q2/q2_execution_detail.csv'); assert len(x)==144288 and set(x.policy)=={'B1','B3','B4'}
def test_q2_shared_g_within_policy_day():
    x=pd.read_csv(ROOT/'outputs/q2/q2_day_ahead_plans.csv'); e=pd.read_csv(ROOT/'outputs/q2/q2_execution_detail.csv'); m=x.merge(e,on=['policy','date','slot']); assert np.allclose(m.grid_plan_kwh,m.grid_purchase_kwh)
def test_cvar_hand_case():
    cfg={'soc_min_kwh':0,'soc_max_kwh':0,'max_charge_power_kw':0,'max_discharge_power_kw':0,'charge_efficiency':.9,'discharge_efficiency':.9}; b3=solve_two_stage_dispatch([[0],[10]],[[0],[0]],[1],[.9,.1],0,cfg,.95,0); b40=solve_two_stage_dispatch([[0],[10]],[[0],[0]],[1],[.9,.1],0,cfg,.95,0); r=solve_two_stage_dispatch([[0],[10]],[[0],[0]],[1],[.9,.1],0,cfg,.95,.2); assert np.isclose(b3.grid_purchase_kwh[0],0) and np.isclose(b3.objective_yuan,5) and np.isclose(b40.objective_yuan,b3.objective_yuan); assert np.isclose(r.grid_purchase_kwh[0],10) and np.isclose(r.objective_yuan,10); assert np.isclose(empirical_cvar([0,50],[.9,.1],.95),50)
def test_result2_readback():
    w=load_workbook(ROOT/'outputs/q2/result2.xlsx',data_only=True); ws=w['计划购电量']; assert ws.max_row==335 and ws.max_column==147 and isinstance(ws['B2'].value,(int,float))
def test_future_mutation_does_not_change_prior_forecast():
    d=read_processed(ROOT/'data/processed/actual_10min.csv'); target=pd.Timestamp('2025-02-01'); a=forecast_day(d,target,'ridge','mean7'); d2=d.copy(); d2.loc[d2.date>target,'load_kw']*=99; b=forecast_day(d2,target,'ridge','mean7'); assert np.allclose(a.load_pred_kw,b.load_pred_kw) and np.allclose(a.pv_pred_kw,b.pv_pred_kw)
def test_receding_fixed_g_zero_error_case():
    cfg=load_config()['battery']; g=np.array([1.,1.,1.]);load=np.array([2.,0.,2.]);pv=np.array([0.,3.,0.]);x=execute_fixed_grid_plan('2025-01-01',g,load,pv,load,pv,np.array([1.,1.,10.]),6000,cfg,'case',6000);assert np.isclose(x.emergency_purchase_kwh.sum(),0) and np.isclose(x.soc_end_kwh.iloc[-1],6000)

def test_future_actual_mutation_does_not_change_prior_actions():
    cfg=load_config()['battery']; g=np.ones(6)*2; pred=np.ones(6); actual_load=pred.copy(); actual_pv=np.zeros(6)
    baseline=execute_fixed_grid_plan('2025-01-01',g,actual_load,actual_pv,pred,np.zeros(6),np.ones(6),6000,cfg,'base')
    changed_load=actual_load.copy(); changed_pv=actual_pv.copy(); changed_load[3:]=20.; changed_pv[3:]=5.
    changed=execute_fixed_grid_plan('2025-01-01',g,changed_load,changed_pv,pred,np.zeros(6),np.ones(6),6000,cfg,'changed')
    cols=['charge_kwh','discharge_kwh','emergency_purchase_kwh','remaining_energy_kwh']
    assert np.allclose(baseline.loc[:2,cols],changed.loc[:2,cols],atol=1e-9,rtol=0)
def test_receding_feedback_responds_to_current_observation():
    cfg=load_config()['battery'];g=np.ones(3)*2;pred=np.array([1.,1.,1.]);a=execute_fixed_grid_plan('2025-01-01',g,np.array([5.,1.,1.]),np.zeros(3),pred,np.zeros(3),np.ones(3),6000,cfg,'a');b=execute_fixed_grid_plan('2025-01-01',g,np.array([1.,1.,1.]),np.zeros(3),pred,np.zeros(3),np.ones(3),6000,cfg,'b');assert not np.allclose(a.charge_kwh,b.charge_kwh) or not np.allclose(a.discharge_kwh,b.discharge_kwh) or not np.allclose(a.emergency_purchase_kwh,b.emergency_purchase_kwh)
