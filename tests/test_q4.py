from pathlib import Path
import os
import subprocess
import sys

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from microgrid.io import get_daily_price, get_q2_daily_price, read_attachment4_prices
from microgrid.models.deterministic_dispatch import solve_deterministic_dispatch
from microgrid.q4 import slot_from_template_label


ROOT = Path(__file__).parents[1]


def test_attachment4_price_changes_optimizer_decision():
    battery = {
        "soc_min_kwh": 0, "soc_max_kwh": 10,
        "max_charge_power_kw": 60, "max_discharge_power_kw": 60,
        "charge_efficiency": 1, "discharge_efficiency": 1,
    }
    load = np.array([0.0, 10.0]); pv = np.zeros(2)
    cheap_first = solve_deterministic_dispatch(load, pv, [1, 10], 0, None, battery)
    cheap_second = solve_deterministic_dispatch(load, pv, [10, 1], 0, None, battery)
    assert not np.allclose(cheap_first.grid_purchase_kwh, cheap_second.grid_purchase_kwh)
    assert cheap_first.grid_purchase_kwh[0] > 9.9
    assert cheap_second.grid_purchase_kwh[1] > 9.9


def test_explicit_attachment1_switch_reproduces_q2_tariff():
    fixed = get_q2_daily_price(ROOT / "data/raw")
    switched = get_daily_price(ROOT / "data/raw", "2025-06-21", "attachment1")
    assert np.array_equal(fixed.slot.to_numpy(), switched.slot.to_numpy())
    assert np.allclose(fixed.price_yuan_per_kwh, switched.price_yuan_per_kwh, atol=0, rtol=0)


def test_attachment4_time_boundary_and_template_cycle():
    prices = read_attachment4_prices(ROOT / "data/raw")
    assert np.array_equal(prices[prices.date == pd.Timestamp("2025-01-01")].slot, np.arange(1, 145))
    workbook = load_workbook(ROOT / "data/raw/templates/result4-2.xlsx", data_only=True)
    sheet = workbook.worksheets[0]
    assert slot_from_template_label(sheet.cell(1, 2).value) == 2
    assert slot_from_template_label(sheet.cell(1, 145).value) == 1
    assert {slot_from_template_label(sheet.cell(1, col).value) for col in range(2, 146)} == set(range(1, 145))


@pytest.mark.parametrize("field,delta", [
    ("price_yuan_per_kwh", 0.1),
    ("charge_kwh", 1.0),
    ("plan_cost_yuan", 1.0),
])
def test_q4_validator_rejects_tampering(tmp_path, field, delta):
    formal = ROOT / "outputs/q4/q4_2_execution_detail.csv"
    if not formal.exists():
        pytest.skip("Q4 formal run has not been generated yet")
    detail = pd.read_csv(formal)
    detail.loc[0, field] += delta
    bad = tmp_path / f"tampered_{field}.csv"
    detail.to_csv(bad, index=False)
    env = os.environ.copy(); env["Q4_2_DETAIL_PATH"] = str(bad); env["PYTHONPATH"] = str(ROOT / "src")
    result = subprocess.run(
        [sys.executable, "scripts/18_validate_q4.py"], cwd=ROOT, env=env,
        capture_output=True, text=True,
    )
    assert result.returncode != 0
    assert "FAILED CHECKS" in result.stderr or "exception" in result.stderr
