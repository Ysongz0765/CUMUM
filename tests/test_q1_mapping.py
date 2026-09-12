from pathlib import Path
import numpy as np,pandas as pd
from openpyxl import load_workbook
from microgrid.io import read_q1,resolve_raw_file
from microgrid.result_template import build_q1_time_mapping
ROOT=Path(__file__).parents[1]
def test_q1_time_mapping_explicit_cycle():
    q=read_q1(resolve_raw_file(ROOT/'data/raw','attachment1'));m=build_q1_time_mapping(q,ROOT/'data/raw/templates/result1.xlsx');assert len(m)==144 and m.slot.nunique()==144 and m.target_cell.nunique()==144;assert m.loc[m.slot==1,'target_cell'].iloc[0]=='B145';assert m.loc[m.slot==2,'target_cell'].iloc[0]=='B2'
def test_q1_endpoint_interpretation():
    q=read_q1(resolve_raw_file(ROOT/'data/raw','attachment1'));assert q.iloc[0].interval_start==pd.Timestamp('2025-01-01 00:00') and q.iloc[-1].interval_end==pd.Timestamp('2025-01-02 00:00') and q.iloc[-1].source_time_label=='0:00+1'
def test_result1_five_column_readback():
    w=load_workbook(ROOT/'outputs/q1/result1.xlsx',data_only=True);ws=w['充放电量'];assert ws.max_column==5 and ws.cell(2,6).value is None and isinstance(ws['B2'].value,(int,float)) and ws['D2'].value=='0:00'
